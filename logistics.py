"""Weekly logistics guide: the outlets activating in a Monday to Sunday week, plus an Excel export."""
import io
from datetime import date, time, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import text

TODAY_SQL = "(now() AT TIME ZONE 'Africa/Nairobi')::date"
KNOWN_CHANNELS = {"ontrade": "Ontrade", "offtrade": "Offtrade", "tht": "THT"}

# Same column names and order as the Excel guide, without the Elements column
HEADERS = ["#", "Outlet Name", "Physical Location", "Division", "Channel",
           "Kick - off Time", "Date of Activation", "Team Lead Name & Tel", "KDM Name & Tel"]
WIDTHS = [4.5, 31, 36, 13, 13.5, 15, 27, 30, 26]


def _today(engine):
    try:
        with engine.connect() as conn:
            return conn.execute(text(f"SELECT {TODAY_SQL}")).scalar()
    except Exception as e:
        print(f"could not read today's date from the database: {e}")
        return date.today()


def monday_of(d):
    return d - timedelta(days=d.weekday())


def week_label(start, end):
    if start.month == end.month:
        return f"{start.day} \u2013 {end.day} {end:%b}"
    return f"{start.day} {start:%b} \u2013 {end.day} {end:%b}"


def _label(value):
    """'ontrade', 'Off trade' and 'THT ' are tidied the same way as on the dashboard."""
    v = " ".join((value or "").split())
    if not v:
        return ""
    return KNOWN_CHANNELS.get(v.replace(" ", "").lower(), v.title())


def _kickoff(t):
    return f"{t:%I:%M%p}".lstrip("0").lower() if t else ""      # 5:00pm


def _day(d):
    return f"{d:%a} {d.day} {d:%b}"                                # Fri 4 Sep (no %-d: Windows)


def _person(name, phone):
    if name and phone:
        return f"{name} - {phone}"
    return name or phone or ""


def get_guide(engine, week=None):
    """week: any date as text (YYYY-MM-DD); it is moved back to that week's Monday."""
    today = _today(engine)
    try:
        start = monday_of(date.fromisoformat((week or "").strip()))
    except ValueError:
        start = monday_of(today)
    end = start + timedelta(days=6)

    with engine.connect() as conn:
        found = conn.execute(text("""
            SELECT o.id AS outlet_id, o.name, o.physical_loc, o.division, o.channel,
                   a.activation_date, a.kickoff_time,
                   tl.name AS tl_name,   tl.phone AS tl_phone,
                   kdm.name AS kdm_name, kdm.phone AS kdm_phone
            FROM activations a
            JOIN outlets o ON o.id = a.outlet_id
            LEFT JOIN staff tl  ON tl.id  = a.team_lead_id
            LEFT JOIN staff kdm ON kdm.id = a.kdm_id
            WHERE a.activation_date BETWEEN :s AND :e
            ORDER BY o.id, a.activation_date
        """), {"s": start, "e": end}).fetchall()

    # one row per outlet, with every activation day of that week joined together
    outlets = {}
    for r in found:
        o = outlets.setdefault(r.outlet_id, {
            "name": " ".join(r.name.split()), "location": " ".join((r.physical_loc or "").split()),
            "division": _label(r.division), "channel": _label(r.channel),
            "days": [], "time": None, "tl": "", "kdm": "",
        })
        o["days"].append(r.activation_date)
        o["time"] = o["time"] or r.kickoff_time
        o["tl"] = o["tl"] or _person(r.tl_name, r.tl_phone)
        o["kdm"] = o["kdm"] or _person(r.kdm_name, r.kdm_phone)

    ordered = sorted(outlets.values(), key=lambda o: (
        o["division"] == "", o["division"].lower(), o["days"][0], o["time"] or time(23, 59), o["name"].lower()))
    rows = [{
        "n": i, "name": o["name"], "location": o["location"], "division": o["division"],
        "channel": o["channel"], "kickoff": _kickoff(o["time"]),
        "dates": " & ".join(_day(d) for d in o["days"]), "tl": o["tl"], "kdm": o["kdm"],
    } for i, o in enumerate(ordered, 1)]

    return {"start": start, "end": end, "rows": rows, "label": f"{week_label(start, end)} {end.year}",
            "prev": (start - timedelta(days=7)).isoformat(), "next": (start + timedelta(days=7)).isoformat(),
            "current": start == monday_of(today)}


def week_summary(engine):
    """Numbers for the dashboard card: outlets activating this week."""
    start = monday_of(_today(engine))
    end = start + timedelta(days=6)
    try:
        with engine.connect() as conn:
            n = conn.execute(text("SELECT count(DISTINCT outlet_id) FROM activations "
                                  "WHERE activation_date BETWEEN :s AND :e"), {"s": start, "e": end}).scalar()
    except Exception as e:
        print(f"week summary failed: {e}")
        n = 0
    return {"count": n or 0, "label": week_label(start, end)}


def build_xlsx(guide):
    """The guide as an Excel file that looks like the one made by hand."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Logistics guide"

    side = Side(style="thin", color="000000")
    border = Border(left=side, right=side, top=side, bottom=side)
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True)
    head_font = Font(name="Times New Roman", size=12, bold=True, color="000000")
    body_font = Font(name="Times New Roman", size=10, color="000000")
    head_fill = PatternFill("solid", fgColor="E25098")
    division_fill = PatternFill("solid", fgColor="FCE4D6")

    for col, (title, width) in enumerate(zip(HEADERS, WIDTHS), 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font, cell.fill, cell.alignment, cell.border = head_font, head_fill, centre, border
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 25.5

    for i, r in enumerate(guide["rows"], 2):
        values = [r["n"], r["name"], r["location"], r["division"] or None, r["channel"],
                  r["kickoff"], r["dates"], r["tl"] or None, r["kdm"] or None]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=i, column=col, value=value)
            cell.font, cell.alignment, cell.border = body_font, centre, border
        if r["division"]:
            ws.cell(row=i, column=4).fill = division_fill

    # a division is written once and merged down its group, as in the hand-made guide
    first = 2
    rows = guide["rows"]
    for i in range(2, len(rows) + 3):
        same = i - 2 < len(rows) and rows[i - 2]["division"] == rows[first - 2]["division"]
        if not same:
            if rows and rows[first - 2]["division"] and i - 1 > first:
                ws.merge_cells(start_row=first, start_column=4, end_row=i - 1, end_column=4)
            first = i

    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
